# Builds finance_tracker.xlsx.
# Header names are authoritative: they must match SCHEMA in apps-script/00_Config.gs
# exactly, so verifySchema() passes without reconciliation.

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
import datetime as dt

import sys
# Output path; override with: python scripts/build_workbook.py <path>
OUT = sys.argv[1] if len(sys.argv) > 1 else 'finance_tracker.xlsx'

FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
EDIT_FILL = PatternFill('solid', fgColor='FFF2CC')   # yellow = you edit these
NOTE_FONT = Font(name=FONT, size=9, italic=True, color='7F7F7F')
BODY = Font(name=FONT, size=10)
TITLE = Font(name=FONT, bold=True, size=14, color='1F3864')
SECTION = Font(name=FONT, bold=True, size=11, color='1F3864')
BIG = Font(name=FONT, bold=True, size=12)
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '"\u20b9"#,##0.00;[Red]-"\u20b9"#,##0.00'
DATEF = 'yyyy-mm-dd'
DTF = 'yyyy-mm-dd hh:mm:ss'
PCT = '0.0%'

wb = Workbook()
wb.remove(wb.active)


def add_sheet(name, headers, widths=None, freeze='A2'):
    ws = wb.create_sheet(name)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[1].height = 28
    for i, _ in enumerate(headers, start=1):
        w = (widths or {}).get(i, 16)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    return ws


# ----------------------------------------------------------------- Transactions
TXN_HEADERS = [
    'Transaction ID', 'Date', 'Time', 'Account', 'Bank', 'Transaction Type',
    'Amount', 'Currency', 'Merchant', 'Category', 'Subcategory', 'Description',
    'Reference Number', 'Source Email ID', 'Status', 'Notes', 'Raw Text',
    'Imported At', 'Last Modified', 'Excluded From Spending']
tx = add_sheet('Transactions', TXN_HEADERS,
               {1: 26, 2: 12, 3: 10, 4: 12, 5: 8, 6: 16, 7: 14, 8: 9, 9: 26,
                10: 16, 11: 14, 12: 30, 13: 22, 14: 22, 15: 14, 16: 34, 17: 40,
                18: 18, 19: 18, 20: 20})

# One example row. Dated in a past month and flagged excluded, so it cannot
# affect any current-month figure on the Dashboard. Delete it before going live.
example = ['EXAMPLE-DELETE-ME', dt.date(2026, 1, 15), '13:07:15', 'HDFC-CC', 'HDFC',
           'Debit', 968.12, 'INR', 'ZOMATO', 'Food', '', 'A payment was made using your Credit Card',
           '', 'example-not-a-real-gmail-id', 'Example',
           'Example row showing the expected format. Delete before going live.',
           'Rs. 968.12 has been debited ... towards ZOMATO on 15 Jan, 2026',
           dt.datetime(2026, 1, 15, 13, 10), dt.datetime(2026, 1, 15, 13, 10), True]
for i, v in enumerate(example, start=1):
    c = tx.cell(row=2, column=i, value=v)
    c.font = Font(name=FONT, size=10, italic=True, color='808080')
tx.cell(row=2, column=2).number_format = DATEF
tx.cell(row=2, column=7).number_format = MONEY
tx.cell(row=2, column=18).number_format = DTF
tx.cell(row=2, column=19).number_format = DTF

dv_type = DataValidation(
    type='list', formula1='"Debit,Credit,Refund,CC Payment,Self-Transfer"', allow_blank=True)
tx.add_data_validation(dv_type)
dv_type.add('F2:F5000')
dv_status = DataValidation(
    type='list', formula1='"Active,Needs Review,Example"', allow_blank=True)
tx.add_data_validation(dv_status)
dv_status.add('O2:O5000')

# ----------------------------------------------------------------- Accounts
acc = add_sheet('Accounts',
                ['Account ID', 'Bank', 'Account Type', 'Sender Rule',
                 'Last 4 Digits', 'Parser', 'Active'],
                {1: 14, 2: 10, 3: 26, 4: 44, 5: 16, 6: 12, 7: 10})
# Last 4 Digits are placeholders. Replace with your own before use:
#   credit card / (account, debit card) / SBI card.
# Both Sender Rule and Last 4 Digits are comma-separated lists.
ACCOUNTS = [
    ['HDFC-CC', 'HDFC', 'Credit Card', 'alerts@hdfcbank.bank.in', '0000', '', True],
    ['UPI-SAV', 'HDFC', 'Savings / UPI / Debit Card', 'alerts@hdfcbank.bank.in',
     '0000, 1111', '', True],
    ['SBI-CC', 'SBI', 'Credit Card',
     'onlinesbicard@sbicard.com, from@cred.club', '2222', '', True],
]
for r, row in enumerate(ACCOUNTS, start=2):
    for i, v in enumerate(row, start=1):
        c = acc.cell(row=r, column=i, value=v)
        c.font = BODY
        c.border = BOX
        if i in (4, 5, 6):
            c.fill = EDIT_FILL
        if i == 5:
            c.number_format = '@'
            c.alignment = Alignment(horizontal='left')
n = len(ACCOUNTS) + 3
for i, t in enumerate([
    'Sender Rule and Last 4 Digits are comma-separated lists.',
    'HDFC alerts every product (credit card, debit card, UPI, savings) from ONE address, so the '
    'last-4 in the body is what attributes a message \u2014 not the sender.',
    'UPI-SAV therefore carries two numbers: the account number (named in UPI and savings alerts) '
    'and the debit card number (card alerts name only the card).',
    'SBI-CC carries two senders: SBI Card for purchases, CRED for bill payments.',
    'Leave Parser blank. Body sniffing picks correctly; set it only to override an observed misrouting.',
], start=0):
    c = acc.cell(row=n + i, column=1, value=('Note: ' if i == 0 else '') + t)
    c.font = NOTE_FONT
    acc.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=7)

# ----------------------------------------------------------------- Settings
st = add_sheet('Settings', ['Setting', 'Value', 'Notes'], {1: 32, 2: 22, 3: 74})
SETTINGS = [
    ['Gmail Label', 'BankAlerts', 'Label the sync scans.'],
    ['Processing Interval', 15, 'Trigger interval in minutes; snapped to 1/5/10/15/30.'],
    ['Duplicate Review Window (days)', 3, 'Same account + same signed amount inside this window is flagged as a possible duplicate.'],
    ['Default Category', 'Uncategorised', 'Used when no Categories keyword matches.'],
    ['Automation Enabled', False, 'KILL SWITCH. Re-read on every trigger fire. FALSE means the trigger exits immediately.'],
    ['Duplicate Handling', 'flag-and-insert', 'flag-and-insert keeps a suspected duplicate and raises a review row; review-only holds it out of the ledger.'],
    ['Lookback Overlap (days)', 2, 'How far back of already-seen mail each run re-examines. Re-reads are discarded by stable ID.'],
    ['Max Messages Per Run', 100, 'Guard against Apps Script execution timeouts.'],
    ['Base Currency', 'INR', 'Written to the Currency column.'],
    ['Refund Match Window (days)', 45, 'How far back to look for the purchase a refund reverses.'],
    ['Self Transfer VPAs', '', 'Comma-separated list of your own UPI handles. Matches are excluded from spending.'],
    ['Flag Uncategorised For Review', False, 'TRUE raises a review row whenever no keyword matched.'],
]
for r, row in enumerate(SETTINGS, start=2):
    for i, v in enumerate(row, start=1):
        c = st.cell(row=r, column=i, value=v)
        c.font = NOTE_FONT if i == 3 else BODY
        c.border = BOX
        if i == 2:
            c.fill = EDIT_FILL
        if i == 3:
            c.alignment = Alignment(wrap_text=True, vertical='top')

# ----------------------------------------------------------------- Categories
cat = add_sheet('Categories', ['Keyword', 'Category', 'Subcategory'],
                {1: 26, 2: 20, 3: 20})
CATEGORIES = [
    ['ZOMATO', 'Food', 'Delivery'],
    ['SWIGGY', 'Food', 'Delivery'],
    ['BLINKIT', 'Groceries', 'Quick commerce'],
    ['AMAZON PAY', 'Bills', ''],
    ['AMAZON', 'Shopping', ''],
    ['HPCL', 'Fuel', ''],
    ['BPCL', 'Fuel', ''],
    ['ATM WITHDRAWAL', 'Cash', ''],
    ['UBER', 'Transport', ''],
    ['IRCTC', 'Travel', 'Rail'],
]
for r, row in enumerate(CATEGORIES, start=2):
    for i, v in enumerate(row, start=1):
        c = cat.cell(row=r, column=i, value=v)
        c.font = BODY
        c.border = BOX
        c.fill = EDIT_FILL
n = len(CATEGORIES) + 3
for i, t in enumerate([
    'Note: longest matching keyword wins, so AMAZON PAY beats AMAZON. Matching is case-insensitive.',
    'A keyword is matched against the merchant first, then the wider email body \u2014 which is how '
    'ATM WITHDRAWAL catches cash withdrawals, since those alerts carry only a street name as the merchant.',
]):
    c = cat.cell(row=n + i, column=1, value=t)
    c.font = NOTE_FONT
    cat.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=3)

# ----------------------------------------------------------------- Review Queue
rq = add_sheet('Review Queue',
               ['Review ID', 'Date', 'Account', 'Amount', 'Merchant', 'Review Reason',
                'Source Email ID', 'Raw Text', 'Detail', 'Status', 'Created At'],
               {1: 28, 2: 12, 3: 12, 4: 14, 5: 26, 6: 26, 7: 22, 8: 40, 9: 50,
                10: 12, 11: 18})
dv_reason = DataValidation(
    type='list',
    formula1='"Unparsed,Possible Duplicate,Ambiguous Category,Unsupported Bank/Format,Ambiguous Transaction Type"',
    allow_blank=True)
rq.add_data_validation(dv_reason)
dv_reason.add('F2:F2000')
dv_rqstatus = DataValidation(type='list', formula1='"Open,Resolved,Ignored"', allow_blank=True)
rq.add_data_validation(dv_rqstatus)
dv_rqstatus.add('J2:J2000')

# ----------------------------------------------------------------- Audit Log
add_sheet('Audit Log',
          ['Timestamp', 'Transaction ID', 'Field', 'Old Value', 'New Value',
           'Changed By', 'Note'],
          {1: 20, 2: 26, 3: 22, 4: 28, 5: 28, 6: 26, 7: 40})

# ----------------------------------------------------------------- Sync Logs
add_sheet('Sync Logs',
          ['Run ID', 'Started At', 'Finished At', 'Trigger Type', 'Emails Scanned',
           'Imported', 'Duplicates Skipped', 'Sent To Review', 'Parsing Errors',
           'Status', 'Error Message', 'Ignored'],
          {1: 26, 2: 19, 3: 19, 4: 14, 5: 14, 6: 11, 7: 17, 8: 15, 9: 14,
           10: 20, 11: 44, 12: 10})

# ----------------------------------------------------------------- Budget
bud = add_sheet('Budget',
                ['Month', 'Category', 'Budgeted Amount', 'Actual Spend',
                 'Remaining', '% Used'],
                {1: 12, 2: 20, 3: 18, 4: 16, 5: 16, 6: 12})
BUDGET_MONTH = '2026-09'
BUDGET_ROWS = [('Food', 6000), ('Groceries', 5000), ('Fuel', 3000),
               ('Shopping', 4000), ('Bills', 3500), ('Transport', 1500),
               ('Cash', 2000), ('Uncategorised', 1000)]
for r, (category, amount) in enumerate(BUDGET_ROWS, start=2):
    bud.cell(row=r, column=1, value=BUDGET_MONTH).font = BODY
    bud.cell(row=r, column=1).fill = EDIT_FILL
    bud.cell(row=r, column=1).number_format = '@'
    bud.cell(row=r, column=2, value=category).font = BODY
    bud.cell(row=r, column=2).fill = EDIT_FILL
    c = bud.cell(row=r, column=3, value=amount)
    c.font = Font(name=FONT, size=10, color='0000FF')
    c.number_format = MONEY
    c.fill = EDIT_FILL
    # Spend for this month and category, excluding anything flagged out of spending.
    d = bud.cell(row=r, column=4, value=(
        '=SUMPRODUCT((TEXT(Transactions!$B$2:$B$5000,"yyyy-mm")=$A{r})'
        '*(Transactions!$J$2:$J$5000=$B{r})'
        '*(Transactions!$T$2:$T$5000<>TRUE)'
        '*Transactions!$G$2:$G$5000)').format(r=r))
    d.number_format = MONEY
    d.font = BODY
    e = bud.cell(row=r, column=5, value='=C{r}-D{r}'.format(r=r))
    e.number_format = MONEY
    e.font = BODY
    f = bud.cell(row=r, column=6, value='=IFERROR(D{r}/C{r},0)'.format(r=r))
    f.number_format = PCT
    f.font = BODY
    for i in range(1, 7):
        bud.cell(row=r, column=i).border = BOX
n = len(BUDGET_ROWS) + 2
tot = bud.cell(row=n, column=2, value='Total')
tot.font = Font(name=FONT, bold=True, size=10)
for col, letter in ((3, 'C'), (4, 'D'), (5, 'E')):
    c = bud.cell(row=n, column=col,
                 value='=SUM({L}2:{L}{e})'.format(L=letter, e=n - 1))
    c.number_format = MONEY
    c.font = Font(name=FONT, bold=True, size=10)
    c.border = BOX
c = bud.cell(row=n + 2, column=1,
             value='Note: Month is text in yyyy-mm form. Edit the yellow cells; '
                   'Actual Spend, Remaining and % Used are formulas.')
c.font = NOTE_FONT
bud.merge_cells(start_row=n + 2, start_column=1, end_row=n + 2, end_column=6)

# ----------------------------------------------------------------- Dashboard
dash = wb.create_sheet('Dashboard', 0)
dash.sheet_view.showGridLines = False
for col, w in ((1, 4), (2, 30), (3, 18), (4, 16), (5, 16), (6, 16), (7, 12)):
    dash.column_dimensions[get_column_letter(col)].width = w

dash['B2'] = 'Finance Tracker'
dash['B2'].font = TITLE
dash['B3'] = ('Reconstructed from the PRP description \u2014 the original Dashboard was not available. '
              'If you find it, replace this tab only; nothing else depends on it.')
dash['B3'].font = NOTE_FONT
dash.merge_cells('B3:G3')

dash['B5'] = 'This month'
dash['B5'].font = SECTION

rows = [
    ('Month', '=TEXT(TODAY(),"yyyy-mm")', '@'),
    ('Month start', '=DATE(YEAR(TODAY()),MONTH(TODAY()),1)', DATEF),
    ('Month end', '=EOMONTH($C$7,0)', DATEF),   # C7 is the date; C6 is text
    ('Days remaining (incl. today)', '=MAX(0,$C$8-TODAY()+1)', '0'),
]
for i, (label, formula, fmt) in enumerate(rows):
    r = 6 + i
    dash.cell(row=r, column=2, value=label).font = BODY
    c = dash.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    c.font = BODY
    c.border = BOX

dash['B11'] = 'Spending'
dash['B11'].font = SECTION

SPEND = ('=SUMPRODUCT((Transactions!$B$2:$B$5000>=$C$7)'
         '*(Transactions!$B$2:$B$5000<=$C$8)'
         '*(Transactions!$T$2:$T$5000<>TRUE)'
         '*Transactions!$G$2:$G$5000)')
spend_rows = [
    ('Total spend this month', SPEND, MONEY),
    ('Total budget this month', '=SUMIFS(Budget!$C$2:$C$500,Budget!$A$2:$A$500,$C$6)', MONEY),
    ('Remaining', '=$C$13-$C$12', MONEY),
    ('Safe to spend per day', '=IF($C$9<=0,0,MAX(0,$C$14)/$C$9)', MONEY),
    ('Budget used', '=IFERROR($C$12/$C$13,0)', PCT),
]
for i, (label, formula, fmt) in enumerate(spend_rows):
    r = 12 + i
    dash.cell(row=r, column=2, value=label).font = BODY if r != 15 else BIG
    c = dash.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    c.font = BODY if r != 15 else BIG
    c.border = BOX
    if r == 15:
        c.fill = PatternFill('solid', fgColor='E2EFDA')

dash['B18'] = 'Health'
dash['B18'].font = SECTION
health = [
    ('Open review items', "=COUNTIFS('Review Queue'!$J$2:$J$2000,\"Open\")", '0'),
    ('Transactions imported this month',
     '=SUMPRODUCT((Transactions!$B$2:$B$5000>=$C$7)*(Transactions!$B$2:$B$5000<=$C$8)'
     '*(Transactions!$A$2:$A$5000<>""))', '0'),
    ('Last successful sync',
     "=IFERROR(LOOKUP(2,1/('Sync Logs'!$J$2:$J$2000=\"Success\"),'Sync Logs'!$C$2:$C$2000),\"(never)\")",
     DTF),
    ('Last run status',
     "=IFERROR(LOOKUP(2,1/('Sync Logs'!$A$2:$A$2000<>\"\"),'Sync Logs'!$J$2:$J$2000),\"(never run)\")",
     '@'),
    ('Automation enabled',
     "=IFERROR(INDEX(Settings!$B$2:$B$100,MATCH(\"Automation Enabled\",Settings!$A$2:$A$100,0)),\"(not set)\")",
     '@'),
]
for i, (label, formula, fmt) in enumerate(health):
    r = 19 + i
    dash.cell(row=r, column=2, value=label).font = BODY
    c = dash.cell(row=r, column=3, value=formula)
    c.number_format = fmt
    c.font = BODY
    c.border = BOX

dash['B25'] = 'By category (this month)'
dash['B25'].font = SECTION
for i, h in enumerate(['Category', 'Budgeted', 'Spent', 'Remaining', '% Used'], start=2):
    c = dash.cell(row=26, column=i, value=h)
    c.font = HDR_FONT
    c.fill = HDR_FILL
    c.border = BOX
for i in range(len(BUDGET_ROWS)):
    r = 27 + i
    br = 2 + i
    for col, src, fmt in ((2, 'B', '@'), (3, 'C', MONEY), (4, 'D', MONEY),
                          (5, 'E', MONEY), (6, 'F', PCT)):
        c = dash.cell(row=r, column=col, value=(
            '=IFERROR(IF(Budget!$A{br}=$C$6,Budget!${src}{br},""),"")'
        ).format(br=br, src=src))
        c.number_format = fmt
        c.font = BODY
        c.border = BOX

legend_row = 27 + len(BUDGET_ROWS) + 2
dash.cell(row=legend_row, column=2, value='How to use this workbook').font = SECTION
legend = [
    'Yellow cells are yours to edit: Accounts, Settings, Categories, and the Budgeted Amount column on Budget.',
    'Everything else is either written by the sync or calculated. Do not hand-edit Transactions \u2014 '
    'use Finance Sync \u2192 Correct a transaction, which records the before and after in Audit Log.',
    'Transactions row 2 is an example row showing the expected format. It is dated in the past and flagged '
    'Excluded From Spending so it cannot affect any figure here. Delete it before going live.',
    'Safe to spend per day = remaining budget \u00f7 days left in the month, floored at zero.',
    'Spending excludes any row flagged Excluded From Spending \u2014 card bill payments and self-transfers, '
    'which are movements of your own money rather than spending.',
]
for i, t in enumerate(legend):
    c = dash.cell(row=legend_row + 1 + i, column=2, value='\u2022  ' + t)
    c.font = NOTE_FONT
    c.alignment = Alignment(wrap_text=True, vertical='top')
    dash.merge_cells(start_row=legend_row + 1 + i, start_column=2,
                     end_row=legend_row + 1 + i, end_column=7)
    dash.row_dimensions[legend_row + 1 + i].height = 26

wb.save(OUT)
print('wrote', OUT)
print('tabs:', wb.sheetnames)
